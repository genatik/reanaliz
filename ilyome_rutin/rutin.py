# -*- coding: utf-8 -*-
r"""
WES / CES analiz rutini - veri toplama motoru.

olgular.xlsx icindeki Durum sutunu BEKLIYOR olan satirlari isler:
  olguyu platformda (Genomize ya da Ilyome; 'Platform' sutunu) bulur -> isitir
  (warmup) -> tarama eksenlerini calistirir
  -> raporlar\<Olgu Kodu>\ klasorunu acar -> kanit tablolarini (CSV) yazar
  -> satiri VERI_HAZIR yapar.

Yorum, ACMG siniflamasi ve rapor yazimi bu scriptin isi DEGILDIR; onu
Claude yapar. Bu script karar vermez, kanit uretir.

Kullanim:
    BASLAT.bat            (cift tiklama - onerilen)
    python rutin.py
    python rutin.py --liste
    python rutin.py --tamamla "209711-OYK-ZEY-OZC" "CFTR c.2657+5G>A tek alel"
"""
from __future__ import print_function
import argparse
import datetime
import json
import os
import sys
import time
import traceback

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

from openpyxl import load_workbook          # noqa: E402
# Platform istemcileri: ayri kurulumlarda yalniz biri bulunur (Genomize kokunde
# genomize_seq.py, Ilyome kokunde ilyome_seq.py). Eksik olan platformun olgusu
# gelirse acik hata verilir; diger platform etkilenmez.
try:
    import genomize_seq as G                # noqa: E402
    G_HATA = None
except Exception as _e:
    G, G_HATA = None, "%s: %s" % (type(_e).__name__, _e)
try:
    import ilyome_seq as I                  # noqa: E402
    I_HATA = None
except Exception as _e:
    I, I_HATA = None, "%s: %s" % (type(_e).__name__, _e)
import akrabalik as AK                      # noqa: E402
import yollar as Y                          # noqa: E402
try:
    import alphagenome_sorgu as AG          # noqa: E402  (AlphaGenome Atlas; istege bagli)
    AG_HATA = None
except Exception as _e:                     # modul/paket sorunu rutini durdurmasin
    AG, AG_HATA = None, "%s: %s" % (type(_e).__name__, _e)

XLSX = Y.olgular_xlsx()      # Drive gecisinden sonra Drive klasorundedir
RAPOR = Y.rapor_koku()        # olgu klasorleri Drive kokunde (kanit + raporlar)
LOG = os.path.join(HERE, "log")

NL = chr(10)
BASLIK_SATIRI = 3
ISLENECEK = "BEKLIYOR"

SUT = {}   # baslik -> sutun no (1 tabanli)


# --------------------------------------------------------------------------
def log(msg):
    line = "%s  %s" % (datetime.datetime.now().strftime("%H:%M:%S"), msg)
    print(line)
    try:
        with open(os.path.join(LOG, datetime.date.today().isoformat() + ".log"),
                  "a", encoding="utf-8") as f:
            f.write(line + "\n")
    except Exception:
        pass


def excel_ac():
    if not os.path.exists(XLSX):
        sys.exit("olgular.xlsx bulunamadi: %s" % XLSX)
    try:
        wb = load_workbook(XLSX)
    except Exception as e:
        sys.exit("olgular.xlsx acilamadi (%s). Dosya Excel'de acik olabilir." % e)
    ws = wb["OLGULAR"]
    SUT.clear()
    for c in range(1, ws.max_column + 1):
        v = ws.cell(row=BASLIK_SATIRI, column=c).value
        if v:
            SUT[str(v).strip()] = c
    return wb, ws


PENDING = {}   # (satir, sutun) -> deger : bu calismada degistirilen hucreler


def excel_kaydet(wb):
    """Dosyayi TAZE okuyup yalnizca bu calismanin degistirdigi hucreleri yaz.

    Eskiden bellekteki (calisma basinda okunmus) kopya oldugu gibi kaydediliyordu;
    analiz surerken dosyada baska biri/baska bir islem degisiklik yaptiysa o
    degisiklikler sessizce siliniyordu. Artik yalnizca PENDING'deki hucreler
    guncelleniyor, geri kalan disk icerigi korunuyor.
    """
    if not PENDING:
        return True
    # Excel dosyayi acikken tutuyorsa YAZMA. Mac'te openpyxl acik dosyanin
    # uzerine yazabiliyor; kullanicinin o sirada elle girdigi veri (or. HPO)
    # boylece sessizce kayboluyordu. PENDING temizlenmez: bir sonraki
    # excel_kaydet cagrisinda yeniden denenir.
    kilit = os.path.join(os.path.dirname(XLSX), "~$" + os.path.basename(XLSX))
    if os.path.exists(kilit):
        log("!! olgular.xlsx Excel'de ACIK - yazilmadi (elle girdiginiz veri korunsun diye).")
        log("   Dosyayi kapatin; kayit sonraki adimda yeniden denenir.")
        return False
    try:
        taze = load_workbook(XLSX)
        tws = taze["OLGULAR"]
        for (r, c), v in PENDING.items():
            hc = tws.cell(row=r, column=c, value=v)
            if r == BASLIK_SATIRI and c in BASLIK_STIL:
                import copy
                k = tws.cell(row=r, column=BASLIK_STIL[c])
                hc.font, hc.fill = copy.copy(k.font), copy.copy(k.fill)
                hc.border, hc.alignment = copy.copy(k.border), copy.copy(k.alignment)
                tws.column_dimensions[hc.column_letter].width = 24
        taze.save(XLSX)
        PENDING.clear()
        return True
    except PermissionError:
        log("!! olgular.xlsx YAZILAMADI - dosya Excel'de acik. Kapatip tekrar deneyin.")
        return False
    except Exception as e:
        log("!! olgular.xlsx yazilamadi: %s" % e)
        return False


def hucre(ws, r, ad):
    c = SUT.get(ad)
    v = ws.cell(row=r, column=c).value if c else None
    # Google E-Tablolar ile duzenlenen xlsx sayilari ondalikli yazar
    # (215885 -> 215885.0); ornek numarasi "215885.0" olarak aranirsa bulunamaz.
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return v


def yaz(ws, r, ad, deger):
    c = SUT.get(ad)
    if c:
        ws.cell(row=r, column=c, value=deger)   # bellekteki kopya tutarli kalsin
        PENDING[(r, c)] = deger                 # diske yalnizca bu hucre yazilir


# Kurulumun varsayilan platformu yollar.py'den gelir (Genomize koku: yok/GENOMIZE,
# Ilyome koku: PLATFORM = "ILYOME"). Iki sistem AYRI Drive klasorlerinde calisir;
# olgular karismaz. 'Platform' sutunu yalniz Genomize kokunde, istege bagli, satir
# bazinda ayrica tanimlanmak istenirse kullanilir.
VARSAYILAN_PLATFORM = str(getattr(Y, "PLATFORM", "GENOMIZE") or "GENOMIZE").upper()
RUN_SUTUNU = "Ilyome Run" if VARSAYILAN_PLATFORM == "ILYOME" else "Genomize Run"
PLATFORM_SUTUNU = "Platform"      # bos -> varsayilan ; Genomize / Ilyome
ISTEMCI = {}                      # platform -> oturum (her platforma bir kez giris)


def platform_oku(ws, r):
    v = str(hucre(ws, r, PLATFORM_SUTUNU) or "").strip().lower() if PLATFORM_SUTUNU in SUT else ""
    if v.startswith("ily"):
        return "ILYOME"
    if v.startswith("gen"):
        return "GENOMIZE"
    return VARSAYILAN_PLATFORM


def platform_modulu(platform):
    if platform == "ILYOME":
        if I is None:
            raise RuntimeError("ilyome_seq.py yuklenemedi: %s" % I_HATA)
        return I
    if G is None:
        raise RuntimeError("genomize_seq.py yuklenemedi: %s" % G_HATA)
    return G


def platform_adi(platform):
    return "Ilyome" if platform == "ILYOME" else "Genomize"


def istemci(platform):
    """Platform oturumu; ilk kullanimda acilir (parola o anda sorulur/ortamdan okunur)."""
    if platform not in ISTEMCI:
        ISTEMCI[platform] = platform_modulu(platform).Seq()
    return ISTEMCI[platform]


def kimlik_var(platform):
    """Ortamda bu platformun giris bilgisi var mi? (BASLAT yalniz gerekeni sorar.)"""
    if platform == "ILYOME":
        return bool(os.environ.get("ILYOME_USER") and os.environ.get("ILYOME_PASS"))
    return bool(os.environ.get("GENOMIZE_USER") and os.environ.get("GENOMIZE_PASS"))


BASLIK_STIL = {}   # sutun -> stil kaynagi sutunu (yeni baslik hucresi icin)


def sutun_garanti(ws, ad):
    """Baslik satirinda `ad` sutunu yoksa en sona ekle (diske excel_kaydet yazar)."""
    if ad in SUT:
        return SUT[ad]
    c = max(SUT.values()) + 1
    ws.cell(row=BASLIK_SATIRI, column=c, value=ad)
    PENDING[(BASLIK_SATIRI, c)] = ad
    BASLIK_STIL[c] = c - 1
    SUT[ad] = c
    log("  olgular.xlsx: '%s' sutunu eklendi" % ad)
    return c


def run_adi(an):
    r = an.get("run") or {}
    return str(r.get("name") or r.get("id") or "").strip()


def run_metni(kisi_run):
    """{kod: run} -> hucre metni. Hepsi ayni runda ise tek ad."""
    adlar = [v for v in kisi_run.values() if v]
    if not adlar:
        return ""
    if len(set(adlar)) == 1 and len(adlar) == len(kisi_run):
        return adlar[0]
    return "; ".join("%s: %s" % (k, v or "?") for k, v in kisi_run.items())


def run_doldur(ws):
    """Run hucresi bos olan TUM satirlari (durumdan bagimsiz) doldur.
    Isitma gerektirmez; kisi basina tek arama istegi. Giris bilgisi ortamda
    olmayan platformun satirlari atlanir (parola sorusu beklenmedik yerde cikmasin)."""
    dolan = 0
    for r in range(BASLIK_SATIRI + 1, ws.max_row + 1):
        kod = str(hucre(ws, r, "Olgu Kodu") or "").strip()
        if not kod or str(hucre(ws, r, RUN_SUTUNU) or "").strip():
            continue
        platform = platform_oku(ws, r)
        if platform not in ISTEMCI and not kimlik_var(platform):
            continue
        try:
            seq = istemci(platform)
        except (SystemExit, Exception) as e:
            log("  %s oturumu acilamadi, run doldurma atlandi: %s" % (platform_adi(platform), str(e)[:80]))
            continue
        kisiler = kodlar(hucre(ws, r, "Proband No"), hucre(ws, r, "Anne No"),
                         hucre(ws, r, "Baba No"), hucre(ws, r, "Diger Bireyler"))
        kisi_run = {}
        for kk in kisiler:
            try:
                kisi_run[kk] = run_adi(seq.resolve(kk))
            except (SystemExit, Exception) as e:
                kisi_run[kk] = ""
                log("  run bulunamadi: %s / %s (%s)" % (kod, kk, str(e)[:80]))
        metin = run_metni(kisi_run)
        if metin:
            yaz(ws, r, RUN_SUTUNU, metin)
            dolan += 1
    if dolan:
        log("  Run sutunu: %d satir dolduruldu" % dolan)
        excel_kaydet(None)
    return dolan


def kodlar(*degerler):
    """Virgul/bosluk ayrilmis kod listesini duzlestir."""
    out = []
    for d in degerler:
        if not d:
            continue
        for p in str(d).replace(";", ",").split(","):
            p = p.strip()
            if p and p not in out:
                out.append(p)
    return out


# --------------------------------------------------------------------------
HPO_BEKLENIYOR = "HPO_BEKLIYOR"
HPO_BEKLEME_SN = 900


HPO_YOK = ("YOK", "-", "N/A", "NA", "UYGULANMAZ")


def hpo_dolu(ws, r):
    """HPO ya da hedef gen girilmis mi? 'YOK' da bilincli bir giris sayilir."""
    h = str(hucre(ws, r, "HPO (istege bagli)") or "").strip()
    g = str(hucre(ws, r, "Hedef Genler (istege bagli)") or "").strip()
    return bool(h or g)


def hpo_bos_satirlar(ws, satirlar):
    """HPO'su ve hedef geni bos olan satirlar (fenotip ekseni calismayacak olanlar)."""
    return [r for r in satirlar if not hpo_dolu(ws, r)]


def bekleyenler(ws):
    """BEKLIYOR satirlar + HPO'su sonradan yazilmis HPO_BEKLIYOR satirlar.

    (11.09.2026: Mac'teki rutin Windows'ta yazilan HPO'yu 12+ dk goremedi, zaman
    asimina dustu ve satirlar HPO_BEKLIYOR'da kaldi; sonraki calisma onlari almadi.)
    """
    işler = []
    for r in range(BASLIK_SATIRI + 1, ws.max_row + 1):
        durum = (hucre(ws, r, "Durum") or "").strip().upper()
        kod = (hucre(ws, r, "Olgu Kodu") or "").strip()
        if not kod:
            continue
        if durum == ISLENECEK:
            işler.append(r)
        elif durum == HPO_BEKLENIYOR and hpo_dolu(ws, r):
            yaz(ws, r, "Durum", ISLENECEK)
            işler.append(r)
    return işler


def satir_dogrula(ws, r):
    """Zorunlu alanlardan bos olanlari dondur."""
    eksik = []
    for alan in ("Olgu Kodu", "Test", "Analiz Tipi", "Proband No", "Klinik Bilgi"):
        v = hucre(ws, r, alan)
        if v is None or not str(v).strip():
            eksik.append(alan)
    return eksik


def hpo_dogrula(seq, hpo):
    """HPO id'lerini Genomize'in kendi ontolojisine karsi dogrula.

    Genomize'in ontoloji surumu HPO'nun son surumunden eskidir. Bilinmeyen
    bir HP id'si hata vermez; fenotip ekseni sessizce 0 satir doner ve
    "fenotiple iliskili varyant yok" gibi gorunur. Bu bir negatif sonuc
    DEGILDIR - sorgunun hic calismamasidir.

    Endpoint id ile arama desteklemiyorsa (hicbiri dogrulanamiyorsa) liste
    oldugu gibi birakilir; yoksa dogrulanamayanlar ATILIR ve rapora yazilir.
    """
    if not hpo:
        return [], {}
    gecerli, bilinmeyen = [], []
    for h in hpo:
        try:
            bulundu = any((t.get("id") or "").upper() == h.upper()
                          for t in (seq.hpo_search(h) or []))
        except Exception as e:
            log("  HPO dogrulama basarisiz (%s): %s - liste oldugu gibi kullaniliyor" % (h, e))
            return list(hpo), {"durum": "dogrulanamadi"}
        (gecerli if bulundu else bilinmeyen).append(h)

    if not gecerli:
        log("  UYARI: hicbir HPO id'si dogrulanamadi - endpoint id aramasini "
            "desteklemiyor olabilir; liste oldugu gibi kullaniliyor")
        return list(hpo), {"durum": "dogrulanamadi", "denenen": list(hpo)}

    if bilinmeyen:
        log("  UYARI: Genomize ontolojisinde YOK, fenotip ekseninden cikarildi: %s"
            % ", ".join(bilinmeyen))
    return gecerli, ({"durum": "kismi", "bilinmeyen": bilinmeyen} if bilinmeyen
                     else {"durum": "tam"})


def _eksenleri_yaz(seq, an, kk, ham, ozet, kanit, P=None):
    P = P or platform_modulu(VARSAYILAN_PLATFORM)
    """Standart eksen sonuclarini denetle ve kanit CSV'lerine yaz."""
    if getattr(seq, "skipped_axes", None):
        ozet.setdefault("atlanan_eksenler", {})[kk] = dict(seq.skipped_axes)
    # SESSIZ BOS SONUC KONTROLU
    # Motor sicak gorunup (probe 200) tum filtreli sorgular 0 satir
    # donebiliyor. Bu bir negatif sonuc DEGIL, veri toplama hatasidir.
    # Yakalanmazsa rapora "temiz negatif" olarak gecer - yanlis negatif.
    toplam_satir = sum(len(v) for v in ham[kk].values())
    if toplam_satir == 0:
        ham_sayi = seq.sanity(an)
        raise RuntimeError(
            "%s (%s): tum tarama eksenleri 0 satir dondu. Filtresiz sorgu %s "
            "kayit veriyor. Bu bir negatif sonuc degil, veri toplama hatasidir - "
            "olgu RAPORLANMAMALIDIR. Analizi tekrar calistirin."
            % (kk, an["sample_name"], ham_sayi))
    for eksen, satirlar in ham[kk].items():
        temiz = [x for x in satirlar if not P.is_artifact(x, 0.02)]
        art = [x for x in satirlar if P.is_artifact(x, 0.02)]
        P.write_csv(os.path.join(kanit, "%s_%s.csv" % (kk, eksen)), temiz)
        if art:
            P.write_csv(os.path.join(kanit, "%s_%s_ARTEFAKT.csv" % (kk, eksen)), art)
        ozet["eksenler"].setdefault(kk, {})[eksen] = {
            "toplam": len(satirlar), "temiz": len(temiz), "artefakt": len(art)}


def qc_doldur(zorla=False):
    """Tamamlanmis olgularin ozet.json'una eksik kalite/kapsama verisini ekle.

    Isinma gerektirmez; birey basina tek dashboard istegi. Platform ozet.json'daki
    "platform" alanindan okunur (yoksa Genomize); giris bilgisi ortamda olmayan
    platformun olgulari atlanir.
    """
    n, atlanan = 0, 0
    for klasor in sorted(os.listdir(RAPOR)):
        yol = os.path.join(RAPOR, klasor, "ozet.json")
        if not os.path.exists(yol):
            continue
        with open(yol, encoding="utf-8") as f:
            ozet = json.load(f)
        degisti = False
        pz = str(ozet.get("platform") or "").lower()
        platform = "ILYOME" if pz.startswith("ily") else ("GENOMIZE" if pz.startswith("gen") else VARSAYILAN_PLATFORM)
        if platform not in ISTEMCI and not kimlik_var(platform):
            continue
        try:
            seq = istemci(platform)
        except (SystemExit, Exception) as e:
            log("  %s oturumu acilamadi, QC atlandi: %s" % (platform_adi(platform), str(e)[:80]))
            continue
        for kk, b in (ozet.get("bireyler") or {}).items():
            if b.get("kapsama") and not zorla:
                atlanan += 1
                continue
            try:
                an = seq.resolve(b.get("ornek") or kk)
                kap, _ = seq.kapsama(an)
            except BaseException as e:
                log("  qc alinamadi: %s / %s (%s)" % (klasor, kk, str(e)[:70]))
                continue
            if kap:
                b["kapsama"] = kap
                degisti = True
                n += 1
        if degisti:
            with open(yol, "w", encoding="utf-8") as f:
                json.dump(ozet, f, indent=1, ensure_ascii=False)
            log("  qc yazildi: %s" % klasor)
    log("QC doldurma bitti: %d birey guncellendi, %d birey zaten vardi." % (n, atlanan))
    return 0


# --------------------------------------------------------------------------
# AlphaGenome Atlas (Google DeepMind) - aday varyantlarin duzenleyici/splicing
# etki skorlari. Kanit uretir, karar vermez; hicbir kosulda olguyu HATA'ya
# dusurmez (API anahtari yoksa / ag yoksa / paket yoksa adim ATLANIR ve
# ozet.json'a nedeni yazilir). Ciktilar kanit/ icine:
#   alphagenome.csv, alphagenome_detay.csv, alphagenome_ozet.json,
#   alphagenome_rapor.txt (rapor yazan Claude icin hazir Turkce bolum)
# --------------------------------------------------------------------------
def alphagenome_dosyalari(kanit, proband):
    """Sorulacak CSV'ler: probandin eksenleri + aile modeli dosyalari."""
    out = []
    for ad in sorted(os.listdir(kanit)):
        a = ad.lower()
        if not a.endswith(".csv") or "artefakt" in a or "_cnv" in a or a.startswith("alphagenome"):
            continue
        if ad.startswith(proband + "_") or a.startswith("denovo_") or a.startswith("ortak_"):
            out.append(os.path.join(kanit, ad))
    return out


def alphagenome_calistir(out, ozet, proband):
    """Olgu klasoru icin AlphaGenome adimi. Donus: ozet['alphagenome'] sozlugu."""
    kanit = os.path.join(out, "kanit")
    sonuc = {"durum": "atlandi", "neden": None}
    if AG is None:
        sonuc["neden"] = "alphagenome_sorgu.py yuklenemedi: %s" % AG_HATA
        ozet["alphagenome"] = sonuc
        log("  AlphaGenome ATLANDI: %s" % sonuc["neden"])
        return sonuc
    if not proband:
        sonuc["neden"] = "proband kodu yok"
        ozet["alphagenome"] = sonuc
        return sonuc
    try:
        genom = ((ozet.get("bireyler") or {}).get(proband) or {}).get("genom") or "hg38"
        dosyalar = alphagenome_dosyalari(kanit, proband)
        adaylar = AG.adaylari_topla(kanit, dosyalar=dosyalar)
        log("  AlphaGenome Atlas: %d aday varyant (%d dosya, %s)" % (len(adaylar), len(dosyalar), genom))
        oz = AG.calistir(kanit, adaylar=adaylar, genome=genom)
        sonuc = {k: oz.get(k) for k in ("durum", "neden", "aday", "atlas", "model", "bulunamadi",
                                        "hata", "sorulamadi", "kategori", "paket", "sure_sn",
                                        "skorlar", "yuksek_varyantlar", "rapor_dosya")}
        if oz.get("durum") == "tamam":
            sonuc["rapor_notu"] = ("Rapor yazarken kanit/alphagenome_rapor.txt'deki bolumu kapsamli analiz "
                                   "raporuna, yontem cumlesini sonuc/resmi rapora ekleyin; ayrinti "
                                   "kanit/alphagenome.csv (arastirma amacli, ACMG kaniti degildir).")
        if oz.get("durum") == "tamam":
            k = oz.get("kategori") or {}
            log("  AlphaGenome bitti: atlas %s, model %s, bulunamadi %s, hata %s; yuksek %s / orta %s / dusuk %s (%s sn)"
                % (oz.get("atlas"), oz.get("model"), oz.get("bulunamadi"), oz.get("hata"),
                   k.get("yuksek"), k.get("orta"), k.get("dusuk"), oz.get("sure_sn")))
            for v in (oz.get("yuksek_varyantlar") or [])[:8]:
                log("    %-24s %-10s AVI=%s  %s" % (v.get("key"), v.get("gene") or "-", v.get("avi"), v.get("yorum")))
        else:
            log("  AlphaGenome %s: %s" % (oz.get("durum"), oz.get("neden")))
    except Exception as e:                    # asla olguyu dusurme
        sonuc = {"durum": "hata", "neden": "%s: %s" % (type(e).__name__, str(e)[:200])}
        log("  !! AlphaGenome adimi hata verdi (olgu etkilenmedi): %s" % sonuc["neden"])
    ozet["alphagenome"] = sonuc
    return sonuc


def alphagenome_ozet_metni(ozet):
    a = ozet.get("alphagenome") or {}
    if a.get("durum") == "tamam":
        k = a.get("kategori") or {}
        return "AlphaGenome yuksek %s / orta %s" % (k.get("yuksek", 0), k.get("orta", 0))
    return "AlphaGenome %s" % (a.get("durum") or "-")


def komut_alphagenome(kodlar):
    """Var olan olgu klasorlerinde yalniz AlphaGenome adimini (yeniden) calistir.

    Genomize'a baglanmaz; kanit/ CSV'leri yeter. 'TUMU' -> kanit klasoru olan tum olgular.
    """
    if kodlar == ["TUMU"]:
        kodlar = [k for k in sorted(os.listdir(RAPOR))
                  if os.path.isdir(os.path.join(RAPOR, k, "kanit")) and not k.startswith(("_", "."))]
    for kod in kodlar:
        out = os.path.join(RAPOR, kod)
        yol = os.path.join(out, "ozet.json")
        if not os.path.exists(yol):
            log("!! %s: ozet.json yok, atlandi" % kod)
            continue
        with open(yol, encoding="utf-8") as f:
            ozet = json.load(f)
        proband = None
        try:
            with open(os.path.join(out, "KLINIK.txt"), encoding="utf-8") as f:
                for s in f:
                    if s.startswith("Proband"):
                        proband = kodlar_ilk(s.split(":", 1)[1])
                        break
        except OSError:
            pass
        proband = proband or kodlar_ilk(",".join(ozet.get("etkilenen") or []))
        log("=== %s  AlphaGenome (proband %s)" % (kod, proband))
        eski = json.dumps(ozet.get("alphagenome"), sort_keys=True)
        alphagenome_calistir(out, ozet, proband)
        if json.dumps(ozet.get("alphagenome"), sort_keys=True) == eski:
            log("  ozet.json degismedi (ayni sonuc), yeniden yazilmadi")   # Drive esitlemesini tetiklemez
            continue
        with open(yol, "w", encoding="utf-8") as f:
            json.dump(ozet, f, indent=1, ensure_ascii=False)
    return 0


def kodlar_ilk(metin):
    k = kodlar(metin)
    return k[0] if k else None


def olgu_isle(seq, ws, r, P=None):
    """P: platform modulu (genomize_seq ya da ilyome_seq); ikisi de ayni arayuzu sunar."""
    kod = str(hucre(ws, r, "Olgu Kodu")).strip()
    platform = platform_oku(ws, r)
    P = P or platform_modulu(platform)
    test = (str(hucre(ws, r, "Test") or "WES")).strip().upper()
    tip = (str(hucre(ws, r, "Analiz Tipi") or "solo")).strip().lower()
    klinik = str(hucre(ws, r, "Klinik Bilgi") or "").strip()
    import re as _re
    hpo = [h.strip() for h in str(hucre(ws, r, "HPO (istege bagli)") or "").replace(";", ",").split(",")
           if _re.match(r"^HP:\d{7}$", h.strip())]
    etkilenen = kodlar(hucre(ws, r, "Etkilenen Bireyler"))
    hedef_gen = [g.strip().upper() for g in
                 str(hucre(ws, r, "Hedef Genler (istege bagli)") or "").replace(";", ",").split(",")
                 if g.strip()]
    notlar = str(hucre(ws, r, "Not") or "").strip()

    proband = kodlar(hucre(ws, r, "Proband No"))
    anne = kodlar(hucre(ws, r, "Anne No"))
    baba = kodlar(hucre(ws, r, "Baba No"))
    diger = kodlar(hucre(ws, r, "Diger Bireyler"))
    tum = kodlar(*(proband + anne + baba + diger))

    log("")
    log("=== %s  (%s, %s, %s)  bireyler: %s" % (kod, test, tip, platform_adi(platform), ", ".join(tum)))

    out = os.path.join(RAPOR, kod)
    kanit = os.path.join(out, "kanit")
    os.makedirs(kanit, exist_ok=True)

    with open(os.path.join(out, "KLINIK.txt"), "w", encoding="utf-8") as f:
        f.write("Olgu Kodu   : %s\n" % kod)
        f.write("Platform    : %s\n" % platform_adi(platform))
        f.write("Test        : %s\n" % test)
        f.write("Analiz Tipi : %s\n" % tip)
        f.write("Proband     : %s\n" % ", ".join(proband))
        f.write("Anne        : %s\n" % ", ".join(anne))
        f.write("Baba        : %s\n" % ", ".join(baba))
        f.write("Diger       : %s\n" % ", ".join(diger))
        f.write("Etkilenen   : %s\n" % ", ".join(etkilenen or proband))
        f.write("Akrabalik   : %s\n" % (hucre(ws, r, "Akrabalik") or ""))
        f.write("HPO         : %s\n" % ", ".join(hpo))
        f.write("Hedef genler: %s\n" % ", ".join(hedef_gen))
        if notlar:
            f.write("\nNOT (olgu listesinden):\n%s\n" % notlar)
        f.write("\nKLINIK BILGI:\n%s\n" % klinik)

    # --- bireyleri coz + isit
    kisiler = {}
    for kk in tum:
        an = seq.resolve(kk)          # birden fazla eslesirse resolve durur
        log("  %-10s -> %s (%s varyant, %s)" % (kk, an["sample_name"],
                                                an.get("numberofmutations"), an.get("genome_version")))
        if not seq.warmup(an):
            raise RuntimeError("warmup tamamlanamadi: %s" % an["sample_name"])
        kisiler[kk] = an

    yaz(ws, r, RUN_SUTUNU, run_metni({kk: run_adi(an) for kk, an in kisiler.items()}))
    ozet_run = {kk: run_adi(an) for kk, an in kisiler.items()}

    hpo, hpo_durum = hpo_dogrula(seq, hpo)
    if hpo_durum.get("bilinmeyen"):
        with open(os.path.join(out, "KLINIK.txt"), "a", encoding="utf-8") as f:
            f.write(NL + "UYARI: Genomize ontolojisinde bulunmayan HPO id'leri "
                    "fenotip ekseninden cikarildi: "
                    + ", ".join(hpo_durum["bilinmeyen"]) + NL)

    ozet = {"olgu": kod, "test": test, "tip": tip, "hpo": hpo,
            "platform": platform_adi(platform),
            "hpo_dogrulama": hpo_durum,
            "etkilenen": etkilenen or proband, "bireyler": {}, "eksenler": {},
            "genomize_run": ozet_run,
            "tarih": datetime.datetime.now().isoformat(timespec="seconds")}

    # --- her birey icin eksenler
    ham = {}
    for kk, an in kisiler.items():
        d = seq.dashboard(an)
        ozet["bireyler"][kk] = {
            "ornek": an["sample_name"], "hash": an["hash"], "genom": an.get("genome_version"),
            "kit": an.get("kit"), "varyant": an.get("numberofmutations"),
            "ortalama_derinlik": d.get("average_depth"), "toplam_anotasyon": d.get("total_annotation"),
            "okuma": an.get("raw_total_read_count"), "on_target": an.get("ontarget_percent"),
            "kapsama": seq.kapsama(an)[0],   # 5x/20x/50x/100x, target covered ...
            "platform": platform_adi(platform),
            "dizileme_tipi": an.get("sequencing_type"), "pipeline": an.get("workers"),
        }
        log("  eksenler: %s" % kk)
        if hedef_gen:
            # HEDEFE YONELIK MOD: yalniz hedef genler. Genom/ekzom geneli
            # eksenler CALISTIRILMAZ (klinisyen istegi, 9-TUR-SOY: "yalniz F8 ve
            # iliskili genler, genomun geri kalaniyla ilgilenme"). Hedef genlerdeki
            # TUM varyantlar (intronik/UTR dahil) tek dosyaya yazilir; artefakt
            # ayrimi yapilmaz (internal_freq sutunu dosyada durur).
            log("  HEDEFE YONELIK analiz - yalniz: %s" % ", ".join(hedef_gen))
            hrows, hbilgi = P.run_hedef(seq, an, hedef_gen)
            if not hrows:
                ham_sayi = seq.sanity(an)
                raise RuntimeError(
                    "%s: hedef gen ekseni 0 satir dondu (filtresiz sorgu %s kayit). Bu bir "
                    "negatif sonuc degil, veri toplama hatasidir - olgu RAPORLANMAMALIDIR."
                    % (kk, ham_sayi))
            P.write_csv(os.path.join(kanit, "%s_hedef_gen.csv" % kk), hrows)
            ozet.setdefault("hedef_gen", {})[kk] = hbilgi
            ozet["mod"] = "hedefe_yonelik"
            ham[kk] = {}
            log("    %d satir; varyantsiz genler: %s%s" % (
                len(hrows), ", ".join(hbilgi["varyantsiz_genler"]) or "-",
                "  !! KESILDI" if hbilgi["kesildi"] else ""))
        else:
            if test == "WGS":
                log("  UYARI: WGS olgusu ama 'Hedef Genler' bos - standart eksenler "
                    "(kodlayan/splays odakli) calisacak; derin intronik tarama YAPILMAZ.")
                ozet.setdefault("uyarilar", []).append("WGS: hedef gen verilmedi")
            ham[kk] = P.run_axes(seq, an, hpo, 0.02, 1.0)
        if not hedef_gen:
            _eksenleri_yaz(seq, an, kk, ham, ozet, kanit, P)

        b = seq.cnv_batch(an)
        if b:
            items, _ = seq.cnvs(an, b)
            rows = P.cnv_rows(items)
            if hedef_gen:
                tum = len(rows)
                rows = [x for x in rows if P.cnv_hedefte(x, hedef_gen)]
                ozet["bireyler"][kk]["cnv_hedef_filtresi"] = (
                    "%d/%d CNV cagrisi hedef genlerde/F8 bolgesinde" % (len(rows), tum))
            P.write_csv(os.path.join(kanit, "%s_cnv.csv" % kk), rows,
                        getattr(P, "ILYOME_CNV_COLUMNS", P.CNV_COLUMNS))
            ozet["bireyler"][kk]["cnv"] = {
                "toplam": len(rows),
                "50kb_uzeri_temiz": len([x for x in rows
                                         if x["length_kb"] >= 50 and not x["cohort_artifact"]])}

    # --- aile modelleri
    etk = [k for k in (etkilenen or proband) if k in ham]
    if len(etk) >= 2:
        log("  ortak varyant analizi: %s" % " + ".join(etk))
        ozet["ortak"] = {}
        for eksen in ham[etk[0]]:
            kumeler = [set(x["key"] for x in ham[k].get(eksen, [])) for k in etk]
            ortak_key = set.intersection(*kumeler) if kumeler else set()
            ortak = [x for x in ham[etk[0]][eksen]
                     if x["key"] in ortak_key and not P.is_artifact(x, 0.02)]
            P.write_csv(os.path.join(kanit, "ORTAK_%s.csv" % eksen), ortak)
            ozet["ortak"][eksen] = len(ortak)

    ebeveyn = [k for k in (anne + baba) if k in ham]
    if proband and proband[0] in ham and ebeveyn:
        p = proband[0]
        log("  de novo adaylari (ebeveynlerde yok): %s" % ", ".join(ebeveyn))
        ozet["de_novo"] = {}
        for eksen in ham[p]:
            eb = set()
            for k in ebeveyn:
                eb |= set(x["key"] for x in ham[k].get(eksen, []))
            dn = [x for x in ham[p][eksen]
                  if x["key"] not in eb and not P.is_artifact(x, 0.02)]
            P.write_csv(os.path.join(kanit, "DENOVO_%s.csv" % eksen), dn)
            ozet["de_novo"][eksen] = len(dn)
        ozet["de_novo_uyari"] = ("Ebeveyn export'unda yokluk de novo KANITI DEGILDIR; "
                                 "kapsama ve filtre farklari yalanci de novo uretir. "
                                 "Populasyon frekansi >0,001 olan bir varyant de novo olamaz.")

    with open(os.path.join(out, "ozet.json"), "w", encoding="utf-8") as f:
        json.dump(ozet, f, indent=1, ensure_ascii=False)

    parcalar = []
    for kk in ozet["eksenler"]:
        e = ozet["eksenler"][kk]
        parcalar.append("%s: fenotip %s / LoF %s / ClinVarP %s" % (
            kk, e.get("fenotip_iliskili", {}).get("temiz", "-"),
            e.get("nadir_lof", {}).get("temiz", "-"),
            e.get("clinvar_patojenik", {}).get("temiz", "-")))
    for kk, hb in (ozet.get("hedef_gen") or {}).items():
        parcalar.append("%s: HEDEFE YONELIK %d satir (%s)" % (
            kk, hb.get("satir", 0), ", ".join(hb.get("varyant_bulunan_genler") or []) or "-"))
    if "ortak" in ozet:
        parcalar.append("ortak fenotip: %s" % ozet["ortak"].get("fenotip_iliskili", "-"))
    if "de_novo" in ozet:
        parcalar.append("de novo aday: %s" % ozet["de_novo"].get("nadir_lof", "-"))

    # --- AILE ILISKISI KONTROLU
    # Trio/duo cikarimlarinin tamami ornek iliskisinin dogrulugana dayanir;
    # ornek karisirsa hicbir eksen hata vermez, rapor sessizce yanlis cikar.
    if (anne or baba) and proband:
        try:
            satirlar, uyarilar = AK.kontrol(out, proband[0],
                                            (anne or [None])[0], (baba or [None])[0])
            ozet["aile_kontrolu"] = {"rapor": satirlar, "uyari": uyarilar}
            for l in satirlar:
                log("  " + l)
            if uyarilar:
                with open(os.path.join(out, "UYARI_AILE_UYUMSUZLUGU.txt"),
                          "w", encoding="utf-8") as f:
                    f.write("AILE ILISKISI UYUMSUZLUGU - RAPOR YAZILMAMALIDIR" + NL)
                    f.write("=" * 60 + NL + NL)
                    f.write(NL.join(satirlar) + NL + NL)
                    f.write("UYARILAR:" + NL)
                    for u in uyarilar:
                        f.write("  - " + u + NL)
                for u in uyarilar:
                    log("  !! AILE UYARISI: " + u)
                raise RuntimeError(
                    "Aile iliskisi uyumsuz: " + "; ".join(uyarilar)[:200] +
                    " -- Ornek kimligi dogrulanmadan bu olgu RAPORLANMAMALIDIR. "
                    "Ayrinti: UYARI_AILE_UYUMSUZLUGU.txt")
        except RuntimeError:
            raise
        except Exception as e:
            log("  aile kontrolu calistirilamadi: %s" % e)
            ozet["aile_kontrolu"] = {"hata": str(e)}

    # --- AlphaGenome Atlas: probandin aday varyantlari (eksenler + DENOVO + ORTAK).
    # Aile kontrolunden SONRA: uyumsuz olguda (HATA) bosa API cagrisi yapilmasin.
    alphagenome_calistir(out, ozet, proband[0] if proband else None)
    parcalar.append(alphagenome_ozet_metni(ozet))

    with open(os.path.join(out, "ozet.json"), "w", encoding="utf-8") as f:
        json.dump(ozet, f, indent=1, ensure_ascii=False)

    return out, " | ".join(parcalar)


# --------------------------------------------------------------------------
def komut_liste(ws):
    print("\n%-24s %-6s %-8s %-9s %-12s %s" % ("OLGU", "TEST", "TIP", "PLATFORM", "DURUM", "PROBAND"))
    print("-" * 78)
    n = 0
    for r in range(BASLIK_SATIRI + 1, ws.max_row + 1):
        kod = hucre(ws, r, "Olgu Kodu")
        if not kod:
            continue
        d = (hucre(ws, r, "Durum") or "").strip().upper()
        print("%-24s %-6s %-8s %-9s %-12s %s" % (
            str(kod)[:24], hucre(ws, r, "Test") or "", hucre(ws, r, "Analiz Tipi") or "",
            platform_adi(platform_oku(ws, r)), d, hucre(ws, r, "Proband No") or ""))
        if d == ISLENECEK:
            n += 1
    print("-" * 78)
    print("%d olgu BEKLIYOR durumunda.\n" % n)


def komut_tamamla(wb, ws, kod, ozet_metni):
    for r in range(BASLIK_SATIRI + 1, ws.max_row + 1):
        if str(hucre(ws, r, "Olgu Kodu") or "").strip() == kod:
            yaz(ws, r, "Durum", "TAMAMLANDI")
            yaz(ws, r, "Tamamlanma", datetime.datetime.now().strftime("%d.%m.%Y %H:%M"))
            if ozet_metni:
                yaz(ws, r, "Sonuc Ozeti", ozet_metni)
            if excel_kaydet(wb):
                log("TAMAMLANDI olarak isaretlendi: %s" % kod)
            else:
                log("!! %s isaretlenemedi - Excel'i kapatip tekrar calistirin." % kod)
                sys.exit(2)
            return
    log("Olgu bulunamadi: %s" % kod)


def main():
    ap = argparse.ArgumentParser(description="WES/CES analiz rutini")
    ap.add_argument("--liste", action="store_true", help="olgulari listele, calistirma")
    ap.add_argument("--tamamla", nargs=2, metavar=("OLGU", "OZET"),
                    help="bir olguyu TAMAMLANDI isaretle")
    ap.add_argument("--run-doldur", action="store_true",
                    help="yalnizca bos 'Genomize Run' hucrelerini doldur")
    ap.add_argument("--qc-doldur", action="store_true",
                    help="eski olgularin ozet.json'una kalite/kapsama verisini ekle")
    ap.add_argument("--platformlar", action="store_true",
                    help="BEKLIYOR olgularin platformlarini yaz (BASLAT yalniz gerekli parolayi sorar)")
    ap.add_argument("--alphagenome", nargs="+", metavar="OLGU",
                    help="var olan olgu(lar) icin yalniz AlphaGenome Atlas adimini calistir "
                         "(Genomize gerekmez; 'TUMU' = kanit klasoru olan tum olgular)")
    a = ap.parse_args()

    os.makedirs(RAPOR, exist_ok=True)
    os.makedirs(LOG, exist_ok=True)
    if a.alphagenome:
        return komut_alphagenome(a.alphagenome)
    wb, ws = excel_ac()

    if a.platformlar:
        pl = sorted({platform_oku(ws, r) for r in range(BASLIK_SATIRI + 1, ws.max_row + 1)
                     if str(hucre(ws, r, "Olgu Kodu") or "").strip()
                     and str(hucre(ws, r, "Durum") or "").strip().upper() in (ISLENECEK, HPO_BEKLENIYOR)})
        print(" ".join(pl))
        return 0
    if a.liste:
        return komut_liste(ws)
    if a.tamamla:
        return komut_tamamla(wb, ws, a.tamamla[0], a.tamamla[1])

    # --qc-doldur yalnizca tamamlanmis olgularin ozet.json'una kalite
    # verisi ekler; Excel satirlariyla ya da HPO ile isi yoktur. Bu
    # denetim eskiden HPO el sikismasindan SONRA geldigi icin BASLAT'in
    # QC adimi da 15 dk HPO bekliyor, kullanici ayni soruyu iki kez
    # goruyordu (once QC adiminda, sonra asil taramada).
    if a.qc_doldur:
        return qc_doldur()

    # --- HPO'yu ONCE klinik bilgiden otomatik uret (hpo_otomatik.py).
    # Eslesen olgu el sikismaya hic girmez; eslesmeyen olgu eskisi gibi
    # sorulur. Asagidaki bekleyenler() "HPO dolu -> BEKLIYOR" kuralini isletir.
    # --qc-doldur adiminda calismaz (o adim Excel'e dokunmaz).
    if not a.qc_doldur:
        try:
            import hpo_otomatik as _HO
            _yazilan = []
            for _r in range(BASLIK_SATIRI + 1, ws.max_row + 1):
                _kod = str(hucre(ws, _r, "Olgu Kodu") or "").strip()
                if not _kod or hpo_dolu(ws, _r):
                    continue
                if str(hucre(ws, _r, "Durum") or "").strip().upper() not in (
                        ISLENECEK, HPO_BEKLENIYOR):
                    continue
                _t = _HO.terimler(str(hucre(ws, _r, "Klinik Bilgi") or ""))
                if _t:
                    yaz(ws, _r, "HPO (istege bagli)", ", ".join(_t))
                    _yazilan.append((_kod, _t))
            if _yazilan:
                print("")
                print("  Klinik bilgiden HPO uretildi (hpo_otomatik.py):")
                for _kod, _t in _yazilan:
                    log("HPO otomatik uretildi: %s -> %s" % (_kod, ", ".join(_t)))
                excel_kaydet(wb)
        except Exception as _e:
            log("!! HPO otomatik uretimi atlandi: %s" % _e)

    sutun_garanti(ws, RUN_SUTUNU)
    işler = bekleyenler(ws)

    # HPO_BEKLIYOR'da kalmis ve HPO'su HALA bos satirlar hicbir listeye girmez.
    # Eskiden sessizce atlanirlardi: kullanici olgusunun neden islenmedigini
    # goremez, bunun yerine baska bir olgunun analiz edildigini gorurdu.
    unutulan = [str(hucre(ws, r, "Olgu Kodu")).strip()
                for r in range(BASLIK_SATIRI + 1, ws.max_row + 1)
                if str(hucre(ws, r, "Olgu Kodu") or "").strip()
                and str(hucre(ws, r, "Durum") or "").strip().upper() == HPO_BEKLENIYOR
                and not hpo_dolu(ws, r)]
    if unutulan:
        print("")
        print("  !! HPO sutunu hala bos - BU OLGULAR ANALIZ EDILMIYOR: %s" % ", ".join(unutulan))
        print("     HPO'yu Excel'e elle yazmak yerine (rutin dosyayi yeniden")
        print("     yazarken elle yapilan degisiklik kaybolabilir) sunu kullanin:")
        print("       python hpo_yaz.py \"<Olgu Kodu>\" \"HP:0000717, HP:0000750\"")
        log("!! HPO_BEKLIYOR + HPO bos, atlandi: %s" % ", ".join(unutulan))

    # HPO EL SIKISMASI (tasarim: HPO'yu klinik bilgiden Claude planlar).
    # HPO'su bos olgu HPO_BEKLIYOR yapilir; Claude'un izleme gorevi bunu gorur,
    # HPO'yu dogrulayip yazar ve satiri yeniden BEKLIYOR yapar. Rutin burada
    # bekler; HPO gelince taze Excel'le devam eder. Hedefe yonelik olgularda
    # HPO gerekmez. Claude oturumu kapaliysa HPO_BEKLEME_SN sonra kullaniciya sorulur.
    hpo_bos = hpo_bos_satirlar(ws, işler)
    if hpo_bos:
        kodlar_ = [str(hucre(ws, r, "Olgu Kodu")).strip() for r in hpo_bos]
        for r in hpo_bos:
            yaz(ws, r, "Durum", HPO_BEKLENIYOR)
        excel_kaydet(wb)
        print("")
        print("  HPO sutunu bos: %s" % ", ".join(kodlar_))
        print("  Claude klinik bilgiden HPO'lari planlayip yaziyor; rutin bekliyor")
        print("  (en fazla %d dk, genellikle 3-5 dk)..." % (HPO_BEKLEME_SN // 60))
        log("HPO bekleniyor (HPO_BEKLIYOR): %s" % ", ".join(kodlar_))
        t0 = time.time()
        while True:
            time.sleep(20)
            asildi = time.time() - t0 > HPO_BEKLEME_SN
            try:
                wb, ws = excel_ac()
            except SystemExit:
                # Dosya o an yaziliyor olabilir - kisa sureli hata normaldir.
                # Ama Excel'de ACIK kaldiysa bu durum surer ve zaman asimi
                # denetimi bu dala hic ugramadigi icin rutin saatlerce beklerdi
                # (28.09.2026 kaydi: 15 dk yerine 62 dk).
                if not asildi:
                    continue
                print("")
                print("  !! olgular.xlsx %d dk boyunca okunamadi - Excel'de acik olabilir."
                      % (HPO_BEKLEME_SN // 60))
                print("     Dosyayi kapatip rutini yeniden baslatin.")
                log("HPO beklemesi: olgular.xlsx okunamadi (dosya acik?), beklemeden cikildi.")
                break
            kalan = [str(hucre(ws, r, "Olgu Kodu")).strip()
                     for r in range(BASLIK_SATIRI + 1, ws.max_row + 1)
                     if str(hucre(ws, r, "Olgu Kodu") or "").strip() in kodlar_
                     and not hpo_dolu(ws, r)]
            if not kalan:
                log("HPO'lar yazildi, devam ediliyor.")
                break
            if asildi:
                print("")
                print("  !! %d dk icinde HPO yazilmadi: %s" % (HPO_BEKLEME_SN // 60, ", ".join(kalan)))
                print("     (Claude oturumu acik degil olabilir.) Bu olgularda FENOTIP EKSENI CALISMAZ.")
                cevap = ""
                if sys.stdin is not None and sys.stdin.isatty():
                    try:
                        cevap = input("     HPO'suz devam edilsin mi? (e/H): ").strip().lower()
                    except EOFError:
                        cevap = ""
                for r in range(BASLIK_SATIRI + 1, ws.max_row + 1):
                    if str(hucre(ws, r, "Olgu Kodu") or "").strip() in kalan:
                        yaz(ws, r, "Durum", ISLENECEK if cevap in ("e", "evet", "y", "yes")
                            else HPO_BEKLENIYOR)
                excel_kaydet(wb)
                log("HPO beklemesi zaman asimi: %s -> %s" % (
                    ", ".join(kalan), "HPO'suz devam" if cevap in ("e", "evet") else "bekletildi"))
                break
            print("     ... bekleniyor: %s (%.0f sn)" % (", ".join(kalan), time.time() - t0))
        wb, ws = excel_ac()
        sutun_garanti(ws, RUN_SUTUNU)
        işler = bekleyenler(ws)
    bos_run = [r for r in range(BASLIK_SATIRI + 1, ws.max_row + 1)
               if str(hucre(ws, r, "Olgu Kodu") or "").strip()
               and not str(hucre(ws, r, RUN_SUTUNU) or "").strip()]
    if not işler and not bos_run and not a.run_doldur:
        excel_kaydet(wb)
        print("\nBEKLIYOR durumunda olgu yok. Yapacak bir sey yok.\n")
        komut_liste(ws)
        return

    if bos_run or a.run_doldur:
        run_doldur(ws)
    if a.run_doldur or not işler:
        excel_kaydet(wb)
        if not işler:
            print("\nBEKLIYOR durumunda olgu yok.\n")
        return

    log("%d olgu islenecek." % len(işler))

    basarili, hatali = [], []
    for r in işler:
        kod = str(hucre(ws, r, "Olgu Kodu")).strip()
        eksik = satir_dogrula(ws, r)
        if eksik:
            yaz(ws, r, "Durum", "HATA")
            yaz(ws, r, "Sonuc Ozeti", "Eksik alan: " + ", ".join(eksik))
            hatali.append(kod)
            log("!! %s atlandi - eksik alan: %s" % (kod, ", ".join(eksik)))
            continue
        try:
            platform = platform_oku(ws, r)
            seq = istemci(platform)                 # ilk olguda giris yapilir
            klasor, ozet = olgu_isle(seq, ws, r, platform_modulu(platform))
            yaz(ws, r, "Durum", "VERI_HAZIR")
            yaz(ws, r, "Rapor Klasoru", os.path.relpath(klasor, RAPOR))
            yaz(ws, r, "Sonuc Ozeti", ozet)
            basarili.append(kod)
            log("   -> %s" % klasor)
        except SystemExit as e:
            yaz(ws, r, "Durum", "HATA")
            yaz(ws, r, "Sonuc Ozeti", str(e)[:250])
            hatali.append(kod)
            log("!! %s HATA: %s" % (kod, e))
        except Exception as e:
            yaz(ws, r, "Durum", "HATA")
            yaz(ws, r, "Sonuc Ozeti", ("%s: %s" % (type(e).__name__, e))[:250])
            hatali.append(kod)
            log("!! %s HATA: %s" % (kod, e))
            traceback.print_exc()
        excel_kaydet(wb)

    print("")
    print("=" * 70)
    print("  VERI TOPLAMA BITTI")
    print("=" * 70)
    print("  Hazir : %s" % (", ".join(basarili) or "-"))
    print("  Hatali: %s" % (", ".join(hatali) or "-"))
    print("")
    try:
        import drive_esitle
        drive_esitle.esitle(sessiz=True)
    except Exception as e:
        log("!! Drive esitleme calismadi: %s" % e)
    print("")
    print("  SIRADAKI ADIM: Claude raporlari KENDILIGINDEN yazar (izleme gorevi")
    print("  VERI_HAZIR olgulari gorur). Windows'taki Claude oturumu acik olmali.")
    print("=" * 70)


if __name__ == "__main__":
    main()
