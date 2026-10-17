# -*- coding: utf-8 -*-
r"""olgular_olustur.py - Ilyome olgu listesini (olgular.xlsx) BOS olarak olusturur.

ILYOME_BASLAT ilk calistirmada olgular.xlsx yoksa bunu kendisi cagirir. Var olan
listeye ASLA dokunmaz (uzerine yazmaz). Genomize listesiyle ayni sutun duzeni,
ayni acilir listeler; yalniz Ilyome olgulari icindir.

    python3 olgular_olustur.py            # yoksa olustur
    python3 olgular_olustur.py --yol X    # baska bir yere (deneme icin)
"""
from __future__ import print_function
import os
import sys

sys.dont_write_bytecode = True
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

SUTUNLAR = [  # (ad, genislik, kullanici_mi)
    ("Olgu Kodu", 22, True), ("Test", 8, True), ("Analiz Tipi", 11, True),
    ("Proband No", 16, True), ("Anne No", 14, True), ("Baba No", 14, True),
    ("Diger Bireyler", 22, True), ("Etkilenen Bireyler", 26, True), ("Akrabalik", 12, True),
    ("Klinik Bilgi", 60, True), ("HPO (istege bagli)", 26, True), ("Oncelik", 10, True),
    ("Durum", 14, True), ("Rapor Klasoru", 30, False), ("Sonuc Ozeti", 46, False),
    ("Tamamlanma", 18, False), ("Not", 40, True), ("Hedef Genler (istege bagli)", 28, True),
    ("Ilyome Run", 24, False), ("Raporlama", 28, True),
]
LISTELER = {"Test": "WES,CES", "Analiz Tipi": "solo,duo,trio,kuartet",
            "Akrabalik": "Var,Yok,Bilinmiyor", "Oncelik": "Normal,Acil",
            "Durum": "BEKLIYOR,TAMAMLANDI,HATA,BEKLEMEDE,ÖRNEK"}
SON_SATIR = 300

KILAVUZ = [
    ("Nasıl çalışır", "1) OLGULAR sayfasına satırı girin. 2) HPO'yu boş bırakabilirsiniz (klinik bilgiden "
     "otomatik üretilir ya da Claude planlar). 3) Durum'u BEKLIYOR yapın. 4) Excel'i KAPATIN. "
     "5) Masaüstündeki ILYOME_BASLAT.command'a çift tıklayın. 6) Raporlar Claude Code varsa kendiliğinden "
     "yazılır; yoksa Claude'a \"Ilyome rutini çalıştı, raporları yaz\" deyin."),
    None,
    ("Bu liste", "YALNIZ Ilyome (ilyome.com) olguları içindir. Genomize olguları "
     "WES_CES_Rutin_Raporlar/olgular.xlsx listesine girilir; iki sistem birbirine karışmaz."),
    ("ZORUNLU alanlar", "Olgu Kodu, Test, Analiz Tipi, Proband No, Klinik Bilgi, Durum"),
    ("Olgu Kodu", "Rapor klasörünün adı olur (Ilyome_Rutin_Raporlar/<Olgu Kodu>/). Hasta adı yerine anonim kod önerilir."),
    ("Proband No", "Ilyome'daki ÖRNEK ADI (sampleid). Birden çok örnek eşleşirse rutin durur ve tam adı ister — "
     "yanlış örneği analiz etmez."),
    ("Anne / Baba No", "Yalnız trio (veya kuartet) için; Ilyome örnek adları."),
    ("Diğer Bireyler", "Duo'da etkilenen kardeş, kuartette ikinci kardeş vb. Birden fazlaysa virgülle ayırın."),
    ("Etkilenen Bireyler", "KRİTİK ALAN. Klinik olarak etkilenmiş bireyler. İki etkilenen kardeşte ortak varyant "
     "analizi yapılır; trio'da etkilenmemiş ebeveynler de novo analizinde kullanılır."),
    ("Klinik Bilgi", "Yaş, cinsiyet, bulgular, akrabalık, aile öyküsü, önceki testler, ön tanı. Fenotip olmadan hangi "
     "bulgunun nedensel olduğu söylenemez."),
    ("HPO", "Boş bırakılabilir: klinik bilgiden otomatik üretilir; üretilemezse rutin HPO_BEKLIYOR'a geçer ve Claude "
     "yazar. HPO hiç olmazsa fenotip ekseni ATLANIR."),
    ("Hedef Genler", "Doluysa YALNIZ bu genler incelenir (intronik/UTR dahil); ekzom geneli eksenler çalışmaz."),
    None,
    ("Durum değerleri", None),
    ("  BEKLIYOR", "Analiz edilecek."),
    ("  VERI_HAZIR", "Veri toplandı; rapor bekliyor."),
    ("  TAMAMLANDI", "Raporlar hazır. BİR DAHA ANALİZ EDİLMEZ."),
    ("  HATA", "Analiz yapılamadı; Sonuç Özeti'nde nedeni yazar."),
    ("  BEKLEMEDE", "Siz beklettiniz (ör. klinik bilgi eksik). Atlanır."),
    ("  ÖRNEK", "Örnek satır. Atlanır."),
    None,
    ("Üretilen raporlar", "İndeks olgu: kapsamlı analiz raporu, sonuç raporu, EÜTF resmi raporu; trio/duo'da her "
     "ebeveyn için resmi rapor. Ayrıca kanit/ (tarama tabloları, CSV), KLINIK.txt ve ozet.json."),
    None,
    ("Önemli", "Gri sütunları (Rapor Klasörü, Sonuç Özeti, Tamamlanma, Ilyome Run) elle doldurmayın."),
    ("Önemli", "Raporlar KARAR DESTEĞİDİR; 'karar-desteği taslağı' uyarısıyla üretilir; değerlendirme ve imza genetik "
     "uzmanına aittir."),
    ("Önemli", "Excel dosyası analiz sırasında AÇIK OLMAMALIDIR — açıkken rutin satırı güncelleyemez."),
]


def olustur(yol):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.utils import get_column_letter as L

    teal, gri = PatternFill("solid", fgColor="2E8B8B"), PatternFill("solid", fgColor="808080")
    giris, oto = PatternFill("solid", fgColor="FFF7E6"), PatternFill("solid", fgColor="EEEEEE")
    ince = Side(style="thin", color="D0D0D0")
    kenar = Border(left=ince, right=ince, top=ince, bottom=ince)
    n = len(SUTUNLAR)

    wb = Workbook()
    ws = wb.active
    ws.title = "OLGULAR"
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=17)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=17)
    a1 = ws.cell(1, 1, "ILYOME — WES / CES ANALİZ RUTİNİ — OLGU LİSTESİ (yalnız ilyome.com olguları)")
    a1.font = Font(size=14, bold=True, color="FFFFFF")
    a1.fill = PatternFill("solid", fgColor="1F4E79")
    a1.alignment = Alignment(vertical="center")
    a2 = ws.cell(2, 1, "Turuncu sütunları siz doldurun. Gri sütunlara dokunmayın — analiz bittiğinde otomatik dolar. "
                       "Proband/Anne/Baba No sütunlarına Ilyome'daki ÖRNEK ADINI (sampleid) yazın. Bir olgunun analiz "
                       "edilmesi için Durum BEKLIYOR olmalıdır; TAMAMLANDI olan satırlar bir daha analiz edilmez. "
                       "Genomize olguları bu listeye GİRİLMEZ.")
    a2.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[1].height = 24
    ws.row_dimensions[2].height = 30
    ws.row_dimensions[3].height = 31.5
    for c, (ad, gen, kul) in enumerate(SUTUNLAR, 1):
        h = ws.cell(3, c, ad)
        h.font = Font(bold=True, color="FFFFFF")
        h.fill = teal if kul else gri
        h.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        h.border = kenar
        ws.column_dimensions[L(c)].width = gen
        for r in range(4, SON_SATIR + 1):
            hc = ws.cell(r, c)
            hc.fill = giris if kul else oto
            hc.border = kenar
            if ad in ("Klinik Bilgi", "Sonuc Ozeti", "Not"):
                hc.alignment = Alignment(wrap_text=True, vertical="top")
    sut = {ad: c for c, (ad, _, _) in enumerate(SUTUNLAR, 1)}
    for ad, liste in LISTELER.items():
        dv = DataValidation(type="list", formula1='"%s"' % liste, allow_blank=True)
        dv.add("%s4:%s%d" % (L(sut[ad]), L(sut[ad]), SON_SATIR))
        ws.add_data_validation(dv)
    for kol in ("Proband No", "Anne No", "Baba No", "Diger Bireyler", "Etkilenen Bireyler"):
        for r in range(4, SON_SATIR + 1):
            ws.cell(r, sut[kol]).number_format = "@"     # 0123 gibi kodlar sayiya donmesin
    ornek = {"Olgu Kodu": "ORNEK-001", "Test": "WES", "Analiz Tipi": "trio", "Proband No": "ORNEK-P",
             "Anne No": "ORNEK-A", "Baba No": "ORNEK-B", "Akrabalik": "Yok",
             "Klinik Bilgi": "Örnek satır: 8 yaş erkek, retinitis pigmentosa, gece körlüğü. Silip kendi olgunuzu girin.",
             "Oncelik": "Normal", "Durum": "ÖRNEK", "Not": "Örnek satır — rutin bu satırı atlar"}
    for ad, v in ornek.items():
        ws.cell(4, sut[ad], v)
    ws.freeze_panes = "A4"
    ws.auto_filter.ref = "A3:%s%d" % (L(n), SON_SATIR)

    k = wb.create_sheet("KILAVUZ")
    k.column_dimensions["A"].width = 26
    k.column_dimensions["B"].width = 100
    b = k.cell(1, 2, "ILYOME WES / CES ANALİZ RUTİNİ — KULLANIM")
    b.font = Font(size=14, bold=True)
    r = 3
    for satir in KILAVUZ:
        if satir:
            a, v = satir
            k.cell(r, 1, a).font = Font(bold=True)
            if v:
                k.cell(r, 2, v).alignment = Alignment(wrap_text=True, vertical="top")
        r += 1
    wb.save(yol)
    return yol


def main():
    yol = None
    if "--yol" in sys.argv:
        yol = sys.argv[sys.argv.index("--yol") + 1]
    else:
        import yollar as Y
        yol = Y.olgular_xlsx()
    if os.path.exists(yol):
        print("  olgular.xlsx zaten var, dokunulmadi: %s" % yol)
        return 0
    olustur(yol)
    print("  Bos Ilyome olgu listesi olusturuldu: %s" % yol)
    print("  Excel'de acip olgularinizi girin (Durum = BEKLIYOR), kapatip ILYOME_BASLAT'i yeniden calistirin.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
