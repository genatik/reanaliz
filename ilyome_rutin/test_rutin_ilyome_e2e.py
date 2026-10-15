# -*- coding: utf-8 -*-
"""Uctan uca test: yamali rutin.py + ilyome_seq.py, SAHTE Ilyome sunucusuyla.

Gercek giris yok. requests.Session.request yama ile kesilir ve Ilyome'un demo
yanitlari (test_fixtures/) dondurulur. Boylece Mac'te calistirmadan once
rutin -> ilyome_seq -> kanit CSV -> ozet.json zinciri dogrulanir.

    python3 test_rutin_ilyome_e2e.py
"""
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
FX = os.path.join(HERE, "test_fixtures")

SAHTE_SUNUCU = r'''
import json, os, re, sys
import requests
FX = %r
def fx(n):
    with open(os.path.join(FX, n), encoding="utf-8") as f:
        return json.load(f)
PANEL = fx("snv-filter-panel.json")
LOG = []
class Y:
    def __init__(self, status, body, url):
        self.status_code = status; self._b = body; self.url = url
        self.text = json.dumps(body) if not isinstance(body, str) else body
    def json(self):
        if isinstance(self._b, str): raise ValueError("not json")
        return self._b
def _snv_sayfa(params):
    rows = fx("snv-list.json")
    page = int(params.get("page", 1)); ps = int(params.get("page_size", 200))
    data = rows["data"] if page == 1 else []
    return {"total": len(rows["data"]), "data": data}
def sahte(self, method, url, **kw):
    params = kw.get("params") or {}
    body = kw.get("json")
    LOG.append((method, url, dict(params), body))
    p = re.sub(r"^https://ilyome\.com", "", url)
    if p.endswith("/api/users/login/"):
        if body.get("password") == "dogru":
            return Y(200, {"access": "A", "refresh": "R", "user": {"id": 1, "email": body["email"],
                     "role": "user", "center": "ege"}}, url)
        return Y(401, {"detail": "Invalid credentials"}, url)
    if p.endswith("/api/users/whoami/"): return Y(200, {"id": 1, "email": "x@y"}, url)
    assert p.startswith("/api/c/ege/") or p.startswith("/api/db/"), "merkez on-eki eksik: " + p
    if "/sample-run/samples/create-list/" in p:
        d = fx("samples-create-list.json")
        q = params.get("sampleid")
        data = [s for s in d["data"] if not q or q in (s["sampleid"], s["patient_name"])]
        if q in ("2308-27", "2308-28"):      # anne/baba: ayni ornegi farkli adla
            s = dict(d["data"][0]); s["id"] = 29446 if q == "2308-27" else 29447
            s["sampleid"] = q; s["patient_name"] = q; data = [s]
        return Y(200, {"total": len(data), "data": data}, url)
    if re.search(r"/sample-run/samples/\d+/overview/snv/top-reportable/", p):
        return Y(200, fx("snv-top-reportable.json"), url)
    if re.search(r"/sample-run/samples/\d+/overview/$", p): return Y(200, fx("samples-overview.json"), url)
    if re.search(r"/ilcore/samples/\d+/d-metrics/", p): return Y(200, fx("metrics-d-metrics.json"), url)
    if re.search(r"/ilcore/samples/\d+/snv/count/", p): return Y(200, {"total": 97966}, url)
    if p.endswith("/ilcore/search-theme/") and method == "GET": return Y(200, {"total": 0, "data": []}, url)
    if p.endswith("/ilcore/search-theme/create/"): return Y(201, {"id": 505, "name": body["name"]}, url)
    if re.search(r"/ilcore/search-theme/505/$", p): return Y(200, PANEL, url)
    if "/ilcore/search-theme/505/reset/" in p: return Y(200, {"detail": "ok"}, url)
    if "/clear-all/" in p: return Y(200, {}, url)
    if "/ilcore/search-theme/update-value/" in p:
        assert method in ("PUT", "PATCH") and set(body) >= {"id", "min", "max", "isnull", "selected", "logic"}
        return Y(200, body, url)
    if "/phenotype/bulk-add/" in p: return Y(200, {"count": len(body["search"].split(","))}, url)
    if re.search(r"/ilcore/search-theme/505/gene/\d+/", p): return Y(200, {}, url)
    if "/api/db/genes/ncbi/" in p:
        return Y(200, {"total": 1, "data": [{"id": 6323, "Symbol": params.get("symbol_i")}]}, url)
    if "/api/db/hpo/search/" in p:
        return Y(200, {"data": [{"term_id": params.get("search"), "name": "Seizure"}]}, url)
    if re.search(r"/ilcore/samples/\d+/snv/505/", p): return Y(200, _snv_sayfa(params), url)
    if re.search(r"/ilcore/samples/\d+/cnv/505/", p):
        page = int(params.get("page", 1)); c = fx("cnv-list.json")
        return Y(200, {"total": len(c["data"]), "data": c["data"] if page == 1 else []}, url)
    return Y(404, {"detail": "sahte sunucu: bilinmeyen uc " + p}, url)
requests.Session.request = sahte
''' % FX


def main():
    tmp = tempfile.mkdtemp(prefix="wes_e2e_")
    drive = os.path.join(tmp, "WES_CES_Rutin_Raporlar")
    sistem = os.path.join(drive, "_sistem")
    os.makedirs(os.path.join(sistem, "log"))
    for f in ("rutin.py", "ilyome_seq.py"):
        shutil.copy(os.path.join(HERE, f), sistem)
    # genomize_seq.py bu repoda yok (Drive'da); COLUMNS icin yoksa ilyome_seq kendi listesini kullanir.
    g = os.path.join(HERE, "genomize_seq.py")
    if os.path.exists(g):
        shutil.copy(g, sistem)
    with open(os.path.join(sistem, "yollar.py"), "w") as f:
        f.write("import os\nK=os.environ['WES_DRIVE_KOKU']\n"
                "def olgular_xlsx(): return os.path.join(K,'olgular.xlsx')\n"
                "def rapor_koku(): return K\n"
                "def sistem_klasoru(): return os.path.join(K,'_sistem')\n")
    with open(os.path.join(sistem, "akrabalik.py"), "w") as f:
        f.write("def kontrol(out, p, a, b): return ['(stub) aile kontrolu'], []\n")
    with open(os.path.join(sistem, "sitecustomize.py"), "w") as f:   # sahte sunucu her python'da yuklensin
        f.write(SAHTE_SUNUCU)
    if not os.path.exists(g):
        with open(os.path.join(sistem, "genomize_seq.py"), "w") as f:  # Genomize istemcisi stub
            f.write("COLUMNS=%r\nCNV_COLUMNS=%r\nHEDEF_PENCERE={}\n"
                    "def cnv_hedefte(r,g): return False\n"
                    "def is_artifact(r,m): return False\n"
                    "def write_csv(*a,**k): pass\n"
                    "class Seq:\n    def __init__(self,*a,**k): raise SystemExit('Genomize cagrilmamali')\n"
                    % (__import__("ilyome_seq").COLUMNS, __import__("ilyome_seq").CNV_COLUMNS))

    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active; ws.title = "OLGULAR"
    ws.append(["OLGULAR"]); ws.append(["aciklama"])
    ws.append(["Olgu Kodu", "Test", "Analiz Tipi", "Proband No", "Anne No", "Baba No", "Diger Bireyler",
               "Etkilenen Bireyler", "Akrabalik", "Klinik Bilgi", "HPO (istege bagli)", "Oncelik", "Durum",
               "Rapor Klasoru", "Sonuc Ozeti", "Tamamlanma", "Not", "Hedef Genler (istege bagli)",
               "Genomize Run", "Platform"])
    ws.append(["2308-TEST-TRIO", "WES", "trio", "2308-26-ZBS", "2308-27", "2308-28", None, None, "Yok",
               "Retinitis pigmentosa, gece korlugu", "HP:0000510, HP:0000662", None, "BEKLIYOR",
               None, None, None, None, None, None, "Ilyome"])
    ws.append(["2308-TEST-HEDEF", "WES", "solo", "2308-26-ZBS", None, None, None, None, "Yok",
               "Hemofili A", "HP:0001892", None, "BEKLIYOR", None, None, None, None, "F8, VWF", None, "Ilyome"])
    wb.save(os.path.join(drive, "olgular.xlsx"))

    env = dict(os.environ, WES_DRIVE_KOKU=drive, ILYOME_USER="test@ege.edu.tr", ILYOME_PASS="dogru",
               PYTHONPATH=sistem, PYTHONDONTWRITEBYTECODE="1")
    r = subprocess.run([sys.executable, "rutin.py"], cwd=sistem, env=env, capture_output=True, text=True,
                       timeout=300)
    out = r.stdout + r.stderr
    print(out[-6000:])
    assert r.returncode == 0, "rutin.py hata kodu %s" % r.returncode

    from openpyxl import load_workbook
    ws = load_workbook(os.path.join(drive, "olgular.xlsx"))["OLGULAR"]
    durumlar = {ws.cell(r, 1).value: ws.cell(r, 13).value for r in (4, 5)}
    print("Durumlar:", durumlar)
    assert durumlar["2308-TEST-TRIO"] == "VERI_HAZIR", durumlar
    assert durumlar["2308-TEST-HEDEF"] == "VERI_HAZIR", durumlar
    assert ws.cell(4, 19).value == "R325-Twist", ws.cell(4, 19).value     # Run sutunu Ilyome run adi

    kanit = os.path.join(drive, "2308-TEST-TRIO", "kanit")
    dosyalar = sorted(os.listdir(kanit))
    print("kanit/:", dosyalar)
    for ad in ("2308-26-ZBS_fenotip_iliskili.csv", "2308-26-ZBS_nadir_lof.csv", "2308-26-ZBS_clinvar_patojenik.csv",
               "2308-26-ZBS_tier.csv", "2308-26-ZBS_homozigot_nadir.csv", "2308-26-ZBS_cnv.csv",
               "2308-27_tier.csv", "DENOVO_tier.csv"):
        assert ad in dosyalar, ad
    with open(os.path.join(kanit, "2308-26-ZBS_fenotip_iliskili.csv"), encoding="utf-8-sig") as f:
        bas = f.readline().strip().split(",")
    assert bas[:4] == ["key", "gene", "gene_mim", "transcript"], bas[:6]
    ozet = json.load(open(os.path.join(drive, "2308-TEST-TRIO", "ozet.json"), encoding="utf-8"))
    assert ozet["platform"] == "Ilyome"
    assert ozet["bireyler"]["2308-26-ZBS"]["genom"] == "hg38"
    assert ozet["bireyler"]["2308-26-ZBS"]["kapsama"], "kapsama bos"
    assert ozet["eksenler"]["2308-26-ZBS"]["nadir_lof"]["toplam"] == 3
    assert ozet["de_novo"]["tier"] == 0           # anne/baba ayni demo verisi -> de novo yok
    print("ozet.json platform/kit/genom:", ozet["platform"], ozet["bireyler"]["2308-26-ZBS"]["kit"],
          ozet["bireyler"]["2308-26-ZBS"]["genom"])
    hedef = json.load(open(os.path.join(drive, "2308-TEST-HEDEF", "ozet.json"), encoding="utf-8"))
    assert hedef.get("mod") == "hedefe_yonelik" and "2308-26-ZBS" in hedef["hedef_gen"]
    assert "2308-26-ZBS_hedef_gen.csv" in os.listdir(os.path.join(drive, "2308-TEST-HEDEF", "kanit"))
    klinik = open(os.path.join(drive, "2308-TEST-TRIO", "KLINIK.txt"), encoding="utf-8").read()
    assert "Platform    : Ilyome" in klinik
    shutil.rmtree(tmp)
    print("\nE2E OK")


if __name__ == "__main__":
    main()
