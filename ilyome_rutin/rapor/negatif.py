# -*- coding: utf-8 -*-
"""Negatif / VUS olgulari icin rapor ureticisi (ILYOME kurulumu).

Genomize surumuyle ayni arayuz; kunye ve kapsama satirlari ozet.json'daki
platform/pipeline alanlarindan ve Ilyome (Dragen) kapsama anahtarlarindan okunur.

Her olgunun yorumu ELLE yazilir; bu modul yalnizca ortak yapiyi (kunye, QC,
aile dogrulamasi, eksen ozeti, kisitliliklar) ozet.json'dan uretir. Karar
vermez.
"""
import io, json, os, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from build_reports import (TPLDIR, NAVY, WARN_TR, setup, banner, para, heading,
                           bullet, table, footer_pagenum, cell_text)
from build7 import (ROOT, TARIH, DOSYA_TARIH, sonuc_raporu, ebeveyn_raporu,
                    TEST_AD, TEST_BASLIK, TEST_KISITLILIK)
from docx.shared import Pt, RGBColor
import shutil

NOTGRUP = "Veri sağlayıcı özetinde bildirilmemiştir (hedef bölge on-target: %99)"

TIP_AD = {"solo": "Solo", "duo": "Duo", "trio": "Trio", "kuartet": "Kuartet"}


def veri(kod):
    kl = os.path.join(ROOT, kod)
    o = json.load(io.open(os.path.join(kl, "ozet.json"), encoding="utf-8"))
    txt = io.open(os.path.join(kl, "KLINIK.txt"), encoding="utf-8").read()
    alan = {}
    for s in txt.splitlines():
        if ":" in s and not s.startswith("KLINIK"):
            a, b = s.split(":", 1)
            alan[a.strip()] = b.strip()
    o["_alan"] = alan
    o["_klinik"] = txt.split("KLINIK BILGI:")[-1].strip()
    o["_klasor"] = kl
    return o


def kunye_satirlari(o):
    rol = [("Proband", "Proband"), ("Anne", "Anne"), ("Baba", "Baba"), ("Diger", "Diğer")]
    sat = [["Olgu Kodu", o["olgu"]],
           ["Analiz tipi", "%s %s + CNV" % (
               TIP_AD.get(o["tip"], o["tip"]),
               TEST_AD.get((o.get("test") or "WES").strip().upper(), TEST_AD["WES"]))]]
    for anahtar, etiket in rol:
        kodlar = [t.strip() for t in (o["_alan"].get(anahtar) or "").split(",") if t.strip()]
        for k in kodlar:
            b = (o.get("bireyler") or {}).get(k)
            if b:
                sat.append([etiket, "%s (%s anotasyon, ortalama derinlik %sx)" % (
                    b.get("ornek"), "{:,}".format(int(b.get("toplam_anotasyon") or 0)).replace(",", "."),
                    round(b.get("ortalama_derinlik") or 0))])
    _kit, _pipe, _plat = None, None, None
    for _b in (o.get("bireyler") or {}).values():
        _kit = _b.get("kit") or _kit
        _pipe = _b.get("pipeline") or _pipe
        _plat = _b.get("platform") or _plat
    _plat = _plat or o.get("platform") or "Ilyome"
    if str(_plat).lower().startswith("ily"):
        _platform_metni = "%s%s — Ilyome (ilyome.com)" % (
            _kit or "WES", (" / %s" % _pipe) if _pipe else "")
    else:
        _platform_metni = "MGI DNBSEQ / %s — Genomize SEQ" % (_kit or "Roche KAPA HyperExome")
    sat += [["Platform / Kit", _platform_metni],
            ["Referans genom", "GRCh38 (hg38)"],
            ["Rapor tarihi", TARIH]]
    return sat


def eksen_cumlesi(o):
    p = [t.strip() for t in (o["_alan"].get("Proband") or "").split(",") if t.strip()][0]
    e = (o.get("eksenler") or {}).get(p) or {}
    parca = []
    for ad, etiket in [("fenotip_iliskili", "fenotip"), ("clinvar_patojenik", "ClinVar-patojenik"),
                       ("nadir_lof", "nadir LoF"), ("tier", "prioritizasyon"),
                       ("homozigot_nadir", "nadir homozigot")]:
        v = e.get(ad)
        if v:
            parca.append("%s ekseni %d" % (etiket, v.get("temiz", 0)))
    atl = o.get("atlanan_eksenler")
    son = "Probandda " + ", ".join(parca) + " temiz kayıt vermiştir; "
    son += ("atlanan eksen yoktur." if not atl else "ATLANAN EKSEN: %s." % atl)
    return son


def metrik(kod, kim=None):
    """Proband (ya da verilen birey) icin resmi rapor metrik satirlari."""
    o = veri(kod)
    if kim is None:
        kim = [t.strip() for t in (o["_alan"].get("Proband") or "").split(",") if t.strip()][0]
    b = (o.get("bireyler") or {}).get(kim) or {}
    ann = int(b.get("toplam_anotasyon") or 0)
    return ["{:,}".format(ann).replace(",", "."),
            "%sx" % round(b.get("ortalama_derinlik") or 0),
            kapsama_metni(b)]


def kapsama_metni(b):
    """Metrik tablosunun '20x kapsayicilik' satiri.

    Genomize dashboard alanlari (rutin --qc-doldur ile ozet.json'a yazilir):
      statistics.target_coverage.total.<N>x.coverage_percent  (1x/5x/20x/50x/100x)
    1x degeri Genomize arayuzundeki "Target Regions Covered" ile aynidir.
    Veri yoksa eskisi gibi 'bildirilmemistir' yazilir.
    """
    kap = b.get("kapsama") or {}

    def yuzde_al(esik):
        for k, v in kap.items():
            a = k.lower()
            # Genomize: statistics.target_coverage.total.20x.coverage_percent
            if "target_coverage" in a and ".total." in a and (".%s." % esik) in a \
                    and a.endswith("coverage_percent"):
                try:
                    return float(v)
                except (TypeError, ValueError):
                    return None
            # Ilyome (Dragen): coverages.PCT of QC coverage region with coverage [  20x: inf)
            if "pct of qc coverage region with coverage" in a and ("%s:" % esik) in a.replace(" ", ""):
                try:
                    return float(v)
                except (TypeError, ValueError):
                    return None
        return None

    def bicim(v):
        return ("%.2f" % v).rstrip("0").rstrip(".").replace(".", ",")

    yirmi, elli, bir = yuzde_al("20x"), yuzde_al("50x"), (yuzde_al("1x") or yuzde_al("3x"))
    if yirmi is not None:
        ek = []
        if elli is not None:
            ek.append("50x: %%%s" % bicim(elli))
        if bir is not None:
            ek.append("hedef bölge kapsaması: %%%s" % bicim(bir))
        return "%%%s" % bicim(yirmi) + ((" (" + "; ".join(ek) + ")") if ek else "")
    if b.get("on_target") is not None:
        return ("Veri sağlayıcı özetinde bildirilmemiştir "
                "(hedef bölge on-target: %%%s)" % b["on_target"])
    return "Veri sağlayıcı özetinde bildirilmemiştir"


def kapsamli(kod, bolumler, klinik_ozet, hpo_metni, aile_tablosu=None,
             aile_yorum=None, qc_uyari=None):
    """bolumler: [(baslik, [paragraf, ...]), ...]"""
    o = veri(kod)
    OUT = o["_klasor"]
    doc = Document(); setup(doc); banner(doc)
    para(doc, TEST_BASLIK.get((o.get("test") or "WES").strip().upper(),
                              TEST_BASLIK["WES"]) + u" — KAPSAMLI ANALİZ RAPORU",
         size=14, bold=True, color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)

    n = 1
    heading(doc, "%d. Künye" % n); n += 1
    table(doc, ["Alan", "Değer"], kunye_satirlari(o), widths=[4.2, 12.3], size=9)

    heading(doc, "%d. Klinik özet ve fenotip kodlaması" % n); n += 1
    para(doc, klinik_ozet)
    if hpo_metni:
        para(doc, "HPO: " + hpo_metni, size=9)
    if qc_uyari:
        para(doc, qc_uyari, size=9, bold=True)

    if aile_tablosu:
        heading(doc, "%d. Örnek kimliği ve aile ilişkisi doğrulaması" % n); n += 1
        table(doc, ["Karşılaştırma", "Ortak nadir varyant", "Oran"], aile_tablosu,
              widths=[6.0, 5.2, 5.3], size=9)
        if aile_yorum:
            para(doc, aile_yorum, size=9)

    for baslik, paragraflar in bolumler:
        heading(doc, "%d. %s" % (n, baslik)); n += 1
        for pr in paragraflar:
            if isinstance(pr, tuple) and pr[0] == "tablo":
                table(doc, pr[1], pr[2], widths=pr[3], size=8.5)
            elif isinstance(pr, tuple) and pr[0] == "madde":
                for m in pr[1]:
                    bullet(doc, m)
            else:
                para(doc, pr)

    heading(doc, "%d. Tarama kapsamı" % n); n += 1
    para(doc, eksen_cumlesi(o), size=9)

    heading(doc, "%d. Kısıtlılıklar" % n)
    para(doc, TEST_KISITLILIK.get((o.get("test") or "WES").strip().upper(),
                                  TEST_KISITLILIK["WES"]) +
              "intronik varyantlar, tekrar genişlemeleri, mitokondriyal genom, dengeli yapısal "
              "değişiklikler, damgalanma bozuklukları ve düşük düzeyli mozaisizm güvenilir biçimde "
              "saptanamaz. Ekzom tabanlı kopya sayısı çağrıları ikinci bir yöntemle "
              "doğrulanmalıdır. Negatif sonuç genetik bir nedeni dışlamaz; yeni gen-hastalık "
              "ilişkileri ortaya çıktıkça yeniden analiz (reanaliz) önerilir.", size=9)

    footer_pagenum(doc, (o.get("test") or "WES").strip().upper()
                        + u" kapsamlı analiz — " + kod)
    p = os.path.join(OUT, "WES_analiz_%s_%s.docx" % (kod, DOSYA_TARIH))
    doc.save(p); return p


def resmi_normal(kod, endikasyon, sonuc_metni, yorum_bloklari, metrikler,
                 tasiyici_satirlari=None):
    """wes_normal.docx sablonu - klinikle iliskili P/LP yok."""
    OUT = os.path.join(ROOT, kod)
    p = os.path.join(OUT, "WES_resmi_%s_%s.docx" % (kod, DOSYA_TARIH))
    shutil.copyfile(os.path.join(TPLDIR, "wes_normal.docx"), p)
    doc = Document(p)
    first = doc.paragraphs[0]
    w = first.insert_paragraph_before()
    r = w.add_run(WARN_TR); r.bold = True; r.font.size = Pt(9)
    r.font.name = "Arial"; r.font.color.rgb = RGBColor(0xB0, 0x00, 0x00)

    for par in doc.paragraphs:
        s = par.text.strip()
        if s.startswith("Test Endikasyonu/Ön tanı"):
            run = par.add_run(" " + endikasyon); run.font.size = Pt(10)
        elif s.startswith("SONUÇ:") and sonuc_metni:
            for rr in list(par.runs):
                rr._element.getparent().remove(rr._element)
            run = par.add_run("SONUÇ: " + sonuc_metni); run.bold = True; run.font.size = Pt(10)

    if tasiyici_satirlari and doc.tables:
        t = doc.tables[0]
        for i, vals in enumerate(tasiyici_satirlari):
            cells = t.rows[i + 1].cells if (i + 1) < len(t.rows) else t.add_row().cells
            for j, v in enumerate(vals):
                if j < len(cells):
                    cell_text(cells[j], v, italic=(j == 0), size=9)

    anchor = None
    for par in doc.paragraphs:
        if par.text.strip() == "Analiz Yorumu:":
            anchor = par; break
    if anchor is not None:
        son = anchor
        for b in yorum_bloklari:
            np = doc.add_paragraph()
            for i, part in enumerate(b.split("*")):
                if not part:
                    continue
                rr = np.add_run(part); rr.font.size = Pt(10); rr.italic = (i % 2 == 1)
            son._p.addnext(np._p); son = np

    for par in list(doc.paragraphs):
        if par.text.strip().startswith("ACMG rehberleri tarafından, ikincil bulgular"):
            par._element.getparent().remove(par._element)

    if doc.tables:
        m = doc.tables[-1]
        for i, v in enumerate(metrikler):
            if i < len(m.rows):
                cell_text(m.rows[i].cells[1], v, size=10)
    doc.save(p); return p
