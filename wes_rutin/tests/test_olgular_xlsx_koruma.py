# -*- coding: utf-8 -*-
"""olgular.xlsx'e elle girilen verinin korunmasi (28.09.2026 duzeltmeleri).

Olay: 215876-BUR-TUR icin HPO sutununa yazilan kodlar her BASLAT'ta kayboldu,
satir HPO_BEKLIYOR'da takildi ve rutin hicbir uyari vermeden baska bir olguyu
(222085-ALM-TOS) analiz etti. Uc ayri kusur vardi:

  1. openpyxl, dosya Excel'de acikken de uzerine yazabiliyor (macOS'ta kilit
     yok) - kullanicinin o sirada girdigi veri siliniyordu.
  2. klinik_dizin.py --doldur, doldurulacak bir sey olmasa da dosyayi her
     calismada bastan yaziyordu.
  3. HPO_BEKLIYOR + HPO'su bos satir hicbir listeye girmedigi icin SESSIZCE
     atlaniyordu.

Testler pipeline dosyalarini (rutin.py, klinik_dizin.py, yollar.py ...) ister;
depo tek basina klonlandiysa atlanirlar.
"""
import contextlib
import io
import json
import os
import shutil
import sys
import tempfile
import time
import unittest

HERE = os.path.dirname(os.path.abspath(__file__))
PIPE = next((p for p in (os.path.dirname(HERE),
                         os.path.join(os.path.dirname(os.path.dirname(HERE)), "genomize_pipeline"),
                         os.environ.get("WES_PIPELINE_DIZINI") or "")
             if p and os.path.exists(os.path.join(p, "rutin.py"))), None)

SUT = ["Olgu Kodu", "Test", "Analiz Tipi", "Proband No", "Anne No", "Baba No",
       "Diger Bireyler", "Etkilenen Bireyler", "Akrabalik", "Klinik Bilgi",
       "HPO (istege bagli)", "Oncelik", "Durum", "Rapor Klasoru", "Sonuc Ozeti",
       "Tamamlanma", "Not", "Hedef Genler (istege bagli)", "Genomize Run", "Raporlama"]

# Gercek dosyanin duzeni: baslik 3. satirda, olgular 4. satirdan itibaren.
SATIRLAR = [
    (4, "215876-BUR-TUR", "HPO_BEKLIYOR", "", "RUN116"),          # HPO bos -> atlanir
    (5, "222085-ALM-TOS", "HPO_BEKLIYOR", "HP:0000717", "RUN132"),  # HPO dolu -> kuyruga
    (6, "218875-ECR-OZE", "TAMAMLANDI", "HP:0011390", "RUN125"),
]


def olgular_uret(p):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "OLGULAR"
    ws.cell(1, 1, "OLGULAR")
    for i, ad in enumerate(SUT, 1):
        ws.cell(3, i, ad)
    S = {ad: i for i, ad in enumerate(SUT, 1)}
    for r, kod, durum, hpo, run in SATIRLAR:
        ws.cell(r, S["Olgu Kodu"], kod)
        ws.cell(r, S["Test"], "WES")
        ws.cell(r, S["Analiz Tipi"], "trio")
        ws.cell(r, S["Proband No"], kod.split("-")[0])
        ws.cell(r, S["Klinik Bilgi"], "test")
        ws.cell(r, S["Durum"], durum)
        ws.cell(r, S["HPO (istege bagli)"], hpo)
        ws.cell(r, S["Genomize Run"], run)
    wb.save(p)


@unittest.skipIf(PIPE is None, "pipeline dosyalari (rutin.py) bulunamadi")
class OlgularKoruma(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.kok = tempfile.mkdtemp(prefix="wes_koruma_")
        os.makedirs(os.path.join(cls.kok, "_sistem"))
        cls.xlsx = os.path.join(cls.kok, "olgular.xlsx")
        olgular_uret(cls.xlsx)
        os.environ["WES_DRIVE_KOKU"] = cls.kok        # yollar.py bu kokU kullanir
        sys.path.insert(0, PIPE)
        sys.dont_write_bytecode = True
        import rutin
        cls.rutin = rutin
        cls.assertEqualPath = rutin.XLSX
        rutin.LOG = os.path.join(cls.kok, "_sistem", "log")
        os.makedirs(rutin.LOG)

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.kok, ignore_errors=True)

    def setUp(self):
        olgular_uret(self.xlsx)
        self.rutin.PENDING.clear()

    def hucre(self, r, ad):
        from openpyxl import load_workbook
        ws = load_workbook(self.xlsx)["OLGULAR"]
        s = {ws.cell(3, c).value: c for c in range(1, ws.max_column + 1) if ws.cell(3, c).value}
        return ws.cell(r, s[ad]).value

    def gunluk(self):
        metin = []
        for a in os.listdir(self.rutin.LOG):
            with io.open(os.path.join(self.rutin.LOG, a), encoding="utf-8") as f:
                metin.append(f.read())
        return "".join(metin)

    # -- 1) Excel acikken yazma --------------------------------------------
    def test_excel_acikken_yazilmaz(self):
        R = self.rutin
        self.assertEqual(R.XLSX, self.xlsx)
        wb, ws = R.excel_ac()
        R.yaz(ws, 4, "Sonuc Ozeti", "TEST")
        kilit = os.path.join(self.kok, "~$olgular.xlsx")
        with io.open(kilit, "w") as f:
            f.write(u"x")
        try:
            self.assertIs(R.excel_kaydet(wb), False)
            self.assertIsNone(self.hucre(4, "Sonuc Ozeti"))
            # kayit kaybolmaz, sonraki denemede yazilir
            self.assertEqual(len(R.PENDING), 1)
        finally:
            os.remove(kilit)
        self.assertIs(R.excel_kaydet(wb), True)
        self.assertEqual(self.hucre(4, "Sonuc Ozeti"), "TEST")
        self.assertEqual(self.hucre(5, "HPO (istege bagli)"), "HP:0000717")

    # -- 2) kuyruga alma ---------------------------------------------------
    def test_hpo_dolunca_kuyruga_alinir_bos_kalirsa_alinmaz(self):
        R = self.rutin
        wb, ws = R.excel_ac()
        kodlar = [str(R.hucre(ws, r, "Olgu Kodu")).strip() for r in R.bekleyenler(ws)]
        self.assertEqual(kodlar, ["222085-ALM-TOS"])

    # -- 3) sessiz atlama --------------------------------------------------
    def test_hpo_bos_satir_sessizce_atlanmaz(self):
        from openpyxl import load_workbook
        R = self.rutin
        wb = load_workbook(self.xlsx)          # kuyrukta is birakma: 222085 -> TAMAMLANDI
        ws = wb["OLGULAR"]
        s = {ws.cell(3, c).value: c for c in range(1, ws.max_column + 1) if ws.cell(3, c).value}
        ws.cell(5, s["Durum"], "TAMAMLANDI")
        wb.save(self.xlsx)
        R.PENDING.clear()

        try:
            from io import StringIO
        except ImportError:                    # py2
            from StringIO import StringIO
        cikti = StringIO()
        sys.argv = ["rutin.py"]
        with contextlib.redirect_stdout(cikti):
            try:
                R.main()
            except SystemExit:
                pass
        metin = cikti.getvalue()
        self.assertIn("215876-BUR-TUR", metin)
        self.assertIn("ANALIZ EDILMIYOR", metin)
        self.assertIn("hpo_yaz.py", metin)     # cozum de gosterilsin
        self.assertIn("HPO_BEKLIYOR + HPO bos, atlandi: 215876-BUR-TUR", self.gunluk())
        self.assertEqual(self.hucre(4, "Durum"), "HPO_BEKLIYOR")
        self.assertEqual(self.hucre(5, "HPO (istege bagli)"), "HP:0000717")

    # -- 4) gereksiz yeniden yazma ----------------------------------------
    def test_klinik_dizin_degisiklik_yokken_dosyaya_dokunmaz(self):
        from openpyxl import load_workbook
        import klinik_dizin
        klinik_dizin.XLSX = self.xlsx
        klinik_dizin.DIZIN = os.path.join(self.kok, "_sistem", "klinik_dizin.json")
        with io.open(klinik_dizin.DIZIN, "w", encoding="utf-8") as f:
            f.write(json.dumps({}))
        try:
            from io import StringIO
        except ImportError:
            from StringIO import StringIO

        zaman, boyut = os.path.getmtime(self.xlsx), os.path.getsize(self.xlsx)
        time.sleep(1.1)                        # mtime cozunurlugu
        with contextlib.redirect_stdout(StringIO()):
            klinik_dizin.komut_doldur()
        self.assertEqual(os.path.getmtime(self.xlsx), zaman)
        self.assertEqual(os.path.getsize(self.xlsx), boyut)

        # ... ama doldurulacak bir sey varsa yine yazar (koruma fazla ileri gitmesin)
        wb = load_workbook(self.xlsx)
        ws = wb["OLGULAR"]
        s = {ws.cell(3, c).value: c for c in range(1, ws.max_column + 1) if ws.cell(3, c).value}
        ws.cell(4, s["Klinik Bilgi"]).value = None   # openpyxl: cell(..., None) deger ATAMAZ
        wb.save(self.xlsx)
        klinik_dizin.ara = lambda kod, dizin: {"klinik": u"otizm, konusma gecikmesi",
                                               "ad_uyumu": True, "klinik_notu_mu": False,
                                               "ad": u"BUR"}
        with contextlib.redirect_stdout(StringIO()):
            klinik_dizin.komut_doldur()
        self.assertEqual(self.hucre(4, "Klinik Bilgi"), u"otizm, konusma gecikmesi")


if __name__ == "__main__":
    unittest.main(verbosity=2)
