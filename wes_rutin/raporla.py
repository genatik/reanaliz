# -*- coding: utf-8 -*-
r"""raporla.py - VERI_HAZIR olgularin rapor surucülerini calistirir ve kapatir.

Zincirin eksik halkasiydi: rutin.py analizi bitirip olguyu VERI_HAZIR yapiyor,
ekrana "Claude raporlari kendiliginden yazar" yaziyordu - ama bunu yapan hicbir
kod yoktu. Raporlar elle uretiliyor, olgu elle --tamamla ile kapatiliyordu.

Bu betik deterministik yarisini yapar:

  _sistem/rapor/vaka_<proband>.py  (ya da vaka_<Olgu Kodu>.py) varsa
    -> calistirir (olgu klasorune .docx dosyalarini yazar)
    -> surucudeki OZET metnini Sonuc Ozeti sutununa yazar
    -> Durum = TAMAMLANDI, Tamamlanma = bugun

Surucu dosyasi YOKSA olgu VERI_HAZIR kalir; yorumu yazacak adim (Claude)
henuz calismamis demektir. Bu betik yorum URETMEZ, karar VERMEZ.

Surucu sozlesmesi (vaka_*.py icinde):
    KOD  = "215876-BUR-TUR"
    OZET = "Negatif. ... (Sonuc Ozeti sutununa yazilacak tek satir)"
    if __name__ == "__main__":  ... raporlari uret ...

    python raporla.py            # bekleyenleri calistir
    python raporla.py --liste    # yalnizca ne yapacagini yaz
"""
from __future__ import print_function
import argparse
import datetime
import os
import re
import subprocess
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
sys.dont_write_bytecode = True

BASLIK = 3
HAZIR = "VERI_HAZIR"
BITTI = "TAMAMLANDI"


def surucu_bul(rapor_dizin, kod, proband):
    """vaka_<proband>.py ya da vaka_<kod>.py; ilk bulunani dondurur."""
    for ad in ("vaka_%s.py" % proband, "vaka_%s.py" % kod):
        p = os.path.join(rapor_dizin, ad)
        if ad and os.path.exists(p):
            return p
    return None


def ozet_oku(surucu):
    """Surucudeki OZET sabitini calistirmadan oku (kaynak metinden)."""
    try:
        with open(surucu, encoding="utf-8") as f:
            kaynak = f.read()
    except Exception:
        return ""
    m = re.search(r'^OZET\s*=\s*(.+?)^\s*$', kaynak, re.S | re.M)
    if not m:
        return ""
    try:
        d = {}
        exec(compile("OZET = " + m.group(1), "<ozet>", "exec"), d)
        v = d.get("OZET")
        return v.strip() if isinstance(v, str) else ""
    except Exception:
        return ""


def main():
    ap = argparse.ArgumentParser(description="Rapor surucülerini calistir ve olguyu kapat")
    ap.add_argument("--liste", action="store_true", help="yalnizca ne yapilacagini yaz")
    ap.add_argument("--olgu", help="yalnizca bu olgu kodu")
    a = ap.parse_args()

    from openpyxl import load_workbook
    import yollar as Y

    p = Y.olgular_xlsx()
    kilit = os.path.join(os.path.dirname(p), "~$" + os.path.basename(p))
    if os.path.exists(kilit):
        print("!! olgular.xlsx Excel'de ACIK. Kapatip tekrar calistirin.")
        return 1
    rapor_dizin = os.path.join(Y.sistem_klasoru(), "rapor")
    if not os.path.isdir(rapor_dizin):
        print("!! rapor klasoru yok: %s" % rapor_dizin)
        return 1

    wb = load_workbook(p)
    ws = wb["OLGULAR"]
    H = {str(ws.cell(BASLIK, c).value).strip(): c
         for c in range(1, ws.max_column + 1) if ws.cell(BASLIK, c).value}

    hazir, surucusuz = [], []
    for r in range(BASLIK + 1, ws.max_row + 1):
        kod = str(ws.cell(r, H["Olgu Kodu"]).value or "").strip()
        if not kod or (a.olgu and kod != a.olgu):
            continue
        if str(ws.cell(r, H["Durum"]).value or "").strip().upper() != HAZIR:
            continue
        pr = ws.cell(r, H["Proband No"]).value
        pr = str(int(pr) if isinstance(pr, float) and pr.is_integer() else pr or "").strip()
        s = surucu_bul(rapor_dizin, kod, pr)
        (hazir if s else surucusuz).append((r, kod, s))

    if not hazir and not surucusuz:
        print("VERI_HAZIR durumunda olgu yok.")
        return 0
    for _, kod, s in hazir:
        print("  + %-22s %s" % (kod, os.path.basename(s)))
    for _, kod, _ in surucusuz:
        print("  - %-22s rapor surucüsü yok (yorum adimi henuz calismamis)" % kod)
    if a.liste:
        print("\n(--liste) hicbir sey calistirilmadi.")
        return 0
    if not hazir:
        print("\nCalistirilacak surucü yok.")
        return 1

    bugun = datetime.date.today().strftime("%d.%m.%Y")
    basarili, basarisiz = [], []
    for r, kod, s in hazir:
        print("\n=== %s" % kod)
        try:
            cik = subprocess.run([sys.executable, os.path.basename(s)], cwd=rapor_dizin,
                                 stdout=subprocess.PIPE, stderr=subprocess.STDOUT,
                                 timeout=900)
            metin = cik.stdout.decode("utf-8", "replace")
        except Exception as e:
            basarisiz.append((kod, str(e)))
            print("  !! calistirilamadi: %s" % e)
            continue
        print("  " + metin.strip().replace("\n", "\n  "))
        if cik.returncode != 0:
            basarisiz.append((kod, "cikis kodu %d" % cik.returncode))
            continue
        ozet = ozet_oku(s)
        ws.cell(r, H["Durum"]).value = BITTI
        if "Tamamlanma" in H:
            ws.cell(r, H["Tamamlanma"]).value = bugun
        if ozet and "Sonuc Ozeti" in H:
            ws.cell(r, H["Sonuc Ozeti"]).value = ozet
        elif not ozet:
            print("  (uyari: surucüde OZET yok - Sonuc Ozeti bos birakildi)")
        if "Rapor Klasoru" in H and not str(ws.cell(r, H["Rapor Klasoru"]).value or "").strip():
            ws.cell(r, H["Rapor Klasoru"]).value = kod
        basarili.append(kod)

    if basarili:
        wb.save(p)
    print("\n%d olgu raporlandi ve TAMAMLANDI yapildi." % len(basarili))
    if basarisiz:
        print("%d olgu basarisiz:" % len(basarisiz))
        for kod, n in basarisiz:
            print("   %-22s %s" % (kod, n))
    if surucusuz:
        print("%d olgu surucü beklemede (VERI_HAZIR kaldi)." % len(surucusuz))
    return 1 if basarisiz else 0


if __name__ == "__main__":
    sys.exit(main())
